from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[2]
out = root / 'output' / 'docx'
out.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.9)
sec.bottom_margin = Inches(.85)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(.4)
sec.footer_distance = Inches(.4)
for name in ['Normal','Title','Heading 1']:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.3
doc.styles['Normal'].paragraph_format.space_after = Pt(10)
doc.styles['Title'].font.size = Pt(16)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(14)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(14)
header = sec.header.paragraphs[0]
header.text = 'DESARROLLO DE UN SISTEMA INTELIGENTE\nDE ANÁLISIS DE RIESGO CREDITICIO'
for r in header.runs:
    r.font.name = 'Arial'
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0,0,0)
foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
f = OxmlElement('w:fldSimple')
f.set(qn('w:instr'),'PAGE')
foot._p.append(f)
sec.different_first_page_header_footer = True

def cover(text, size=12, bold=False, after=12, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

cover('TECNOLÓGICO DE ESTUDIOS SUPERIORES\nDEL ORIENTE DEL ESTADO DE MÉXICO', 15, True, 38)
cover('DESARROLLO DE UN SISTEMA INTELIGENTE\nDE ANÁLISIS DE RIESGO CREDITICIO', 16, True, 30, 'Title')
cover('MEMORIA DE RESIDENCIA PROFESIONAL', 12, True, 8)
cover('INGENIERÍA EN SISTEMAS COMPUTACIONALES', 12, False, 26)
cover('PRESENTA', 11, True, 6)
cover('VILLEGAS RODRIGUEZ JOSE JULIAN', 13, True, 22)
cover('EMPRESA\nPluriOne S.A. de C.V.\nDevelop Talent & Technology', 11, False, 20)
cover('ASESOR EMPRESARIAL\nJuan Méndez Herrera', 11, False, 16)
cover('ASESOR ACADÉMICO\n________________________________', 11, False, 22)
cover('LOS REYES LA PAZ, ESTADO DE MÉXICO\nFECHA DE ENTREGA __________________', 10, False, 0)
doc.add_page_break()
doc.add_paragraph('Índice', 'Title')
entries = [
('1','Introducción'),
('2','Objetivos y justificación'),
('2.1','Objetivo general'),('2.2','Objetivos específicos'),('2.3','Justificación y alcance'),
('3','Marco contextual'),('3.1','Descripción de la empresa'),('3.2','Cultura organizacional y marco normativo'),
('4','Marco conceptual y fundamentos tecnológicos'),
('5','Planeación del proyecto'),('5.1','Acta de constitución del proyecto o Project Charter'),('5.2','Requerimientos funcionales y no funcionales'),('5.3','Estructura de desglose del trabajo o WBS'),('5.4','Cronograma de actividades'),
('6','Metodología Scrum'),('6.1','Roles y backlog del producto'),('6.2','Sprints y ceremonias'),('6.3','Tablero Kanban y seguimiento'),
('7','Diseño de la solución'),('7.1','Arquitectura y diagramas del sistema'),('7.2','Diseño de la base de datos'),('7.3','Prototipo de la interfaz'),
('8','Desarrollo del sistema'),('8.1','Backend y frontend'),('8.2','Modelos predictivos e integraciones'),
('9','Pruebas y resultados'),('10','Mantenimiento y capacitación'),('11','Conclusiones'),('12','Fuentes de información'),('13','Anexos')]
for num,title in entries:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.left_indent = Inches(.22 if '.' in num else 0)
    r = p.add_run(num+'  '+title)
    r.font.size = Pt(10.5)
    r.bold = '.' not in num
doc.add_page_break()
doc.add_paragraph('Introducción', 'Heading 1')
intro = [
'El proyecto de Desarrollo de un Sistema Inteligente de Análisis de Riesgo Crediticio se plantea para PluriOne S.A. de C.V., cuyo nombre comercial es Develop Talent & Technology. Su propósito es construir una aplicación web que apoye la evaluación de solicitudes de crédito mediante el análisis de información financiera y el uso de modelos predictivos. Este trabajo forma parte de la formación profesional en Ingeniería en Sistemas Computacionales y permite aplicar conocimientos de desarrollo de software, bases de datos e inteligencia artificial a una necesidad empresarial.',
'La propuesta contempla analizar historiales crediticios, comportamiento de pago, ingresos, endeudamiento y variables económicas para generar indicadores de riesgo y recomendaciones de aprobación o seguimiento. Se busca que los resultados sean comprensibles para las personas que revisen las solicitudes y que permitan reconocer los factores considerados en cada evaluación. La utilidad de estas recomendaciones deberá comprobarse mediante pruebas y datos adecuados al alcance del proyecto.',
'La solución se desarrollará con Python y FastAPI para el backend, React.js para la interfaz y PostgreSQL para el almacenamiento. Azure Machine Learning se utilizará para los modelos predictivos, mientras que Azure AI Search y Azure OpenAI Service se integrarán para recuperar información y elaborar explicaciones. El alcance tecnológico también incluye Microsoft Entra ID, APIs financieras, Power BI, Docker y GitHub Actions para atender las necesidades de acceso, consulta de datos, reportes y automatización.',
'Actualmente se cuenta con un prototipo local que recibe ingresos, gastos y un score capturado por el usuario. El sistema valida las entradas, calcula el margen mensual disponible y muestra una clasificación preliminar mediante reglas fijas. Esta base permite comprobar la comunicación entre la interfaz y el servidor. El modelo predictivo, la persistencia y los servicios externos corresponden a etapas posteriores de implementación y validación.',
'La organización del trabajo se realizará mediante Scrum, con un backlog de actividades y entregas incrementales. La memoria documentará la planeación, los requerimientos, la arquitectura y el desarrollo del sistema, junto con las pruebas realizadas y sus resultados. También incorporará el mantenimiento, la capacitación y las conclusiones conforme se completen esas actividades. De esta manera, el documento permitirá relacionar los objetivos del proyecto con los avances y las evidencias obtenidas durante su desarrollo.'
]
for text in intro:
    p=doc.add_paragraph(text)
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
doc.core_properties.title='Memoria de residencia del sistema de análisis de riesgo crediticio'
doc.core_properties.author='Villegas Rodriguez Jose Julian'
dest=out/'PluriOne_portada_indice_introduccion.docx'
doc.save(dest)
print(dest)
